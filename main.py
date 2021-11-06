import argparse
from reverseCombiner import reverseCombiner
from openpyxl import load_workbook
from openpyxl.styles import Font

input_file = "raw.xlsx"
output_file = "new_" + input_file

# Open workbook
# workbook = load_workbook(filename="raw.xlsx", read_only=True)
workbook = load_workbook(filename=input_file)

# Get worksheets object, worksheet count for iterator
all_sheets = workbook.worksheets
sheet_count = len(all_sheets)

# Get reference sheet and length for delete algorithm
reference_sheet = workbook.worksheets[3]
reference_sheet_length = len(reference_sheet["A"])

# print(sheet_count)
# print(reference_sheet.title)
# print(reference_sheet_length)

# Get list of rows to delete
rows_list = []

for row in range(reference_sheet_length, 3, -1):
    cell = "A" + str(row)

    if reference_sheet[cell].font.color.rgb != "FFFF0000":
        rows_list.append(row)

rows_tuple = reverseCombiner(rows_list)
print(rows_tuple)

# Loop and delete
for tuple in rows_tuple:
    reference_sheet.delete_rows(tuple[0], tuple[1])

workbook.save(output_file)
